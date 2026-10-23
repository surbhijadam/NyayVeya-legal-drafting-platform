"""
Word document (.docx) builder for Section 138 Legal Notice.
Uses python-docx to generate professionally styled, fully editable legal documents.
"""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE


def build_cheque_bounce_docx(formatted_notice, output_filepath):
    """
    Build and save a clean, professionally formatted .docx legal notice.
    """
    os.makedirs(os.path.dirname(output_filepath), exist_ok=True)

    doc = Document()

    # Set page margins (Standard 1 inch legal layout)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 1. Header (RPAD / Speed Post)
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_header = p_header.add_run(formatted_notice['header'])
    run_header.bold = True
    run_header.font.name = 'Times New Roman'
    run_header.font.size = Pt(11)
    run_header.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Add spacing
    p_header.paragraph_format.space_after = Pt(14)

    # 2. Date
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_date = p_date.add_run(f"Date: {formatted_notice['notice_date']}")
    run_date.font.name = 'Times New Roman'
    run_date.font.size = Pt(11)
    p_date.paragraph_format.space_after = Pt(14)

    # 3. Sender Block (if advocate)
    if 'Advocate' in formatted_notice['sender_block']:
        p_from = doc.add_paragraph()
        p_from.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for line in formatted_notice['sender_block'].split('\n'):
            run = p_from.add_run(line + '\n')
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
        p_from.paragraph_format.space_after = Pt(12)

    # 4. Recipient (TO:) Block
    p_to = doc.add_paragraph()
    p_to.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for line in formatted_notice['recipient_block'].split('\n'):
        run = p_to.add_run(line + '\n')
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        if line.startswith('TO:'):
            run.bold = True
    p_to.paragraph_format.space_after = Pt(16)

    # 5. Subject Line
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run_sub = p_sub.add_run(formatted_notice['subject'])
    run_sub.bold = True
    run_sub.font.name = 'Times New Roman'
    run_sub.font.size = Pt(11)
    p_sub.paragraph_format.space_after = Pt(16)

    # Sir / Madam salutation
    p_sal = doc.add_paragraph()
    run_sal = p_sal.add_run("Sir / Madam,")
    run_sal.font.name = 'Times New Roman'
    run_sal.font.size = Pt(11)
    p_sal.paragraph_format.space_after = Pt(10)

    # 6. Numbered Paragraphs
    for para in formatted_notice['paragraphs']:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(para)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(10)

    # 7. Signoff
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for line in formatted_notice['signoff'].split('\n'):
        run = p_sign.add_run(line + '\n')
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
    p_sign.paragraph_format.space_before = Pt(20)

    doc.save(output_filepath)
    return output_filepath
